"""N1: one WorldCore, real XLSX work and JSON delivery through bound sessions.

The protocol uses literal arithmetic expectations (6 and 8). Neither the runtime
formula engine nor a hidden answer artifact generates the expected values.
This is a finite capability migration experiment, not a professional-quality or
model-capability evaluation. Run only into a fresh output directory.
"""

import argparse
import copy
import io
import json
import traceback
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from proworksim.audit import code_identity
from proworksim.core.world import WorldSpec
from proworksim.storage import atomic_write, digest, json_bytes
from proworksim.world_core import WorldCore

ACTORS = ("alice", "bob", "manager")
INITIAL_MARGIN = 6
DELIVERY_MARGIN = 8
INITIAL_CELLS = {"Report!A1": 10, "Report!A2": 4, "Report!A3": "=A1-A2"}
UPDATED_CELLS = {"Report!A1": 12, "Report!A3": "=SUM(A1,-A2)"}
CHECKS = (
    "one_world_two_projects",
    "same_names_distinct_scoped_identity",
    "initial_workbook_real_bytes_and_scalar",
    "legitimate_edit_new_version_formula_and_scalar",
    "exact_version_sheet_read",
    "recalculate_preserves_formula_and_scalar",
    "json_delivery_actual_bytes",
    "xlsx_submission_pins_current_version",
    "xlsx_fixed_submission_content_passes",
    "json_fixed_submission_content_passes",
    "error_formula_persisted_real_bytes",
    "bad_submission_executes",
    "independent_evaluation_rejects_formula_error",
    "old_submission_and_file_remain_immutable",
    "old_submission_still_passes_after_bad_current",
    "evaluation_reads_only_pinned_versions",
    "evaluation_preserves_state_and_file_bytes",
    "crossproject_samefilename_remains_independent",
    "unauthorized_same_project_edit_rejected",
    "unauthorized_cross_project_edit_rejected",
    "unauthorized_edit_preserves_formal_objects_bytes",
    "common_action_journal_records_both_formats",
)
PROTOCOL = {
    "suite": "work-capability-N1-v0.7",
    "declared_checks": CHECKS,
    "worlds": 1,
    "projects": ["A", "B"],
    "applications": ["files", "spreadsheets"],
    "independent_expected_values": {
        "initial_margin": INITIAL_MARGIN,
        "delivered_margin": DELIVERY_MARGIN,
        "origin": "Literal preregistered scalars: 10 - 4 = 6, 12 - 4 = 8; not computed by UUT",
    },
    "actions": "Install legal packages; create and update XLSX; recalculate; read exact versions; create JSON; submit; independently evaluate fixed immutable bytes; edit to =1/0; attempt unauthorized edits",
    "xlsx_inspection": "openpyxl raw formulas and persisted cached values, without formula evaluation",
    "evaluation_read_comparison": "Exact submitted object/version pairs and SHA256 of their immutable files; no answer-model artifact exists",
    "rejection_comparison": "Artifact registry, immutable bytes and materialized mirrors; rejected attempts may append command diagnostics and advance logical time",
    "scope": "One single-writer WorldCore; two projects share identities, Store and action journal. A owns formula XLSX delivery, B owns JSON delivery plus a same-name XLSX isolation fixture. The script is an informed mechanism driver; public-only worker discovery is N4, not claimed here.",
    "limitations": [
        "Finite formula subset; no complete Excel or financial-quality claim",
        "No model calls, GPU, training, publication/adoption or crash-recovery coverage in this group",
        "A's valid and deliberately invalid submissions belong to separate declared work items",
        "Delivery-only acceptance is an institutional policy result, not independent content correctness",
    ],
}


def write_json(path, value):
    atomic_write(path, json_bytes(value))
    return {"path": str(path), "sha256": digest(path.read_bytes()), "bytes": path.stat().st_size}


def bootstrap(root):
    return WorldCore.create(
        root,
        WorldSpec(
            world_id="work-capability-N1",
            actors={actor: {} for actor in ACTORS},
            applications=["files", "spreadsheets"],
            bootstrap_grants=[
                {"actor_id": "manager", "scope": "world", "power": "install_project"}
            ],
        ),
    )


def package(project_id, owner, kind):
    checks = (
        [
            {
                "kind": "xlsx_cell_equals",
                "sheet": "Report",
                "cell": "A3",
                "expected": DELIVERY_MARGIN,
            },
            {"kind": "xlsx_no_formula_errors"},
        ]
        if kind == "xlsx"
        else [{"kind": "json_field_equals", "path": ["margin"], "expected": DELIVERY_MARGIN}]
    )
    contract = {
        "min_files": 1,
        "max_files": 1,
        "allowed_roles": ["report"],
        "allowed_kinds": [kind],
        "required_fields": [],
        "content_checks": checks,
    }
    work = {
        "work_id": "work-1",
        "owner": owner,
        "goal": "Deliver the public scalar margin 8 in the declared format",
        "approval_policy": "delivery_only",
        "purpose": "delivery",
        "requirement_dimension": "requirements",
        "deliverable_contract": contract,
    }
    works = [work]
    if project_id == "A":
        works.append(
            {
                **copy.deepcopy(work),
                "work_id": "work-error",
                "goal": "Separate controlled erroneous submission with the same content contract",
            }
        )
    return {
        "project_id": project_id,
        "goal": "Finite actual file capability fixture",
        "participants": list(ACTORS),
        "objects": [],
        "works": works,
        "grants": [{"actor_id": owner, "power": "create_object", "subject": "artifact"}],
        "provenance": {
            "kind": "synthetic",
            "source_evidence_refs": [],
            "note": "Independent literal expected values; no professional workflow claim",
        },
        "close_policy": {"pending_obligations": "retain"},
    }


def immutable_bytes(world, reference):
    state = world.store.load()
    artifact = state["artifacts"][reference["object_id"]]
    return world.store.version_path(artifact, reference["version_id"]).read_bytes()


def workbook_cell(content, sheet="Report", cell="A3"):
    """Read actual OOXML formulas and persisted cache, never evaluate formulas."""
    raw = load_workbook(io.BytesIO(content), data_only=False)
    cache = load_workbook(io.BytesIO(content), data_only=True)
    result = {
        "raw": raw[sheet][cell].value,
        "cached": cache[sheet][cell].value,
        "cache_type": cache[sheet][cell].data_type,
    }
    raw.close()
    cache.close()
    return result


def all_file_hashes(world):
    state = world.store.load()
    result = {}
    for aid, artifact in state["artifacts"].items():
        for vid in artifact["versions"]:
            result[f"versions/{aid}/{vid}"] = digest(
                world.store.version_path(artifact, vid).read_bytes()
            )
        result[f"mirrors/{aid}"] = digest(world.store.current_path(artifact).read_bytes())
    return result


def submission(world, project, work, sid):
    return next(
        sub
        for sub in world.store.load()["work_items"][f"{project}::{work}"]["submissions"]
        if sub["submission_id"] == sid
    )


class Evidence:
    def __init__(self, output):
        self.output = output
        self.checks, self.calls, self.snapshots, self.evaluations = [], [], [], []
        self.world = None

    def check(self, name, observed, expected):
        if name not in CHECKS or any(record["name"] == name for record in self.checks):
            raise AssertionError("Undeclared or duplicate check: " + name)
        self.checks.append(
            {
                "name": name,
                "observed": observed,
                "expected": expected,
                "passed": observed == expected,
                "not_executed": False,
            }
        )

    def call(self, session, action, must_succeed=True, **arguments):
        response = session.call(action, **arguments)
        self.calls.append(
            {
                "actor": session.actor_id,
                "project": session.project_id,
                "tool": action,
                "arguments": copy.deepcopy(arguments),
                "response": copy.deepcopy(response),
            }
        )
        if must_succeed and not response.get("ok"):
            raise AssertionError(f"{session.actor_id}/{session.project_id}/{action}: {response}")
        return response["result"] if must_succeed else response

    def snapshot(self, name):
        self.snapshots.append(
            {
                "name": name,
                "state": write_json(self.output / f"{name}-state.json", self.world.store.load()),
                "files": all_file_hashes(self.world),
            }
        )

    def evaluate(self, project, work, sid):
        """Trace immutable path lookups in the evaluator, independently of its read_set."""
        paths = []
        original = self.world.store.version_path

        def traced_path(artifact, version_id):
            paths.append({"object_id": artifact["artifact_id"], "version_id": version_id})
            return original(artifact, version_id)

        self.world.store.version_path = traced_path
        try:
            result = self.world.evaluate_submission(project, work, sid)
        finally:
            self.world.store.version_path = original
        sub = submission(self.world, project, work, sid)
        expected_reads = [
            {
                "object_id": aid,
                "version_id": vid,
                "sha256": digest(
                    original(self.world.store.load()["artifacts"][aid], vid).read_bytes()
                ),
            }
            for aid, vid in sorted(sub["artifact_versions"].items())
        ]
        expected_pairs = [
            {"object_id": ref["object_id"], "version_id": ref["version_id"]}
            for ref in expected_reads
        ]
        record = {
            "project": project,
            "work": work,
            "submission_id": sid,
            "evaluation": result,
            "immutable_path_lookups": paths,
            "expected_reads": expected_reads,
            "exact_read_set": sorted(
                result.get("read_set", []), key=lambda ref: (ref["object_id"], ref["version_id"])
            )
            == expected_reads,
            "exact_path_set": sorted(paths, key=lambda ref: (ref["object_id"], ref["version_id"]))
            == expected_pairs,
        }
        self.evaluations.append(record)
        return result


def run_n1(ev):
    world = ev.world = bootstrap(ev.output / "world")
    initial = world.store.load()
    manager = world.session("manager")
    for pkg in (package("A", "alice", "xlsx"), package("B", "bob", "json")):
        ev.call(manager, "install_project", package=pkg)
    a, b = world.session("alice", "A"), world.session("bob", "B")
    installed = world.store.load()
    ev.check(
        "one_world_two_projects",
        [
            sorted(installed["projects"]),
            installed["world_id"],
            installed["instance_id"] == initial["instance_id"],
            installed["actors"] == initial["actors"],
        ],
        [["A", "B"], "work-capability-N1", True, True],
    )
    first = ev.call(
        a,
        "create_object",
        alias="report",
        filename="report.xlsx",
        kind="xlsx",
        data=INITIAL_CELLS,
        deliverable_role="report",
    )
    b_book = ev.call(
        b,
        "create_object",
        alias="report",
        filename="report.xlsx",
        kind="xlsx",
        data={"Report!A3": 99},
        deliverable_role="draft",
    )
    state = world.store.load()
    ev.check(
        "same_names_distinct_scoped_identity",
        [
            first["object_id"] != b_book["object_id"],
            state["artifacts"][first["object_id"]]["filename"],
            state["artifacts"][b_book["object_id"]]["filename"],
            "A::work-1" in state["work_items"],
            "B::work-1" in state["work_items"],
        ],
        [True, "report.xlsx", "report.xlsx", True, True],
    )
    first_bytes = immutable_bytes(world, first)
    ev.check(
        "initial_workbook_real_bytes_and_scalar",
        [first_bytes.startswith(b"PK"), workbook_cell(first_bytes)],
        [True, {"raw": "=A1-A2", "cached": INITIAL_MARGIN, "cache_type": "n"}],
    )
    updated = ev.call(a, "sheet_update", alias="report", cells=UPDATED_CELLS)
    updated_bytes = immutable_bytes(world, updated)
    ev.check(
        "legitimate_edit_new_version_formula_and_scalar",
        [
            updated["object_id"] == first["object_id"],
            updated["version_id"] != first["version_id"],
            digest(updated_bytes) != digest(first_bytes),
            workbook_cell(updated_bytes),
        ],
        [True, True, True, {"raw": "=SUM(A1,-A2)", "cached": DELIVERY_MARGIN, "cache_type": "n"}],
    )
    old_read = ev.call(a, "sheet_read", alias="report", version_id=first["version_id"])
    new_read = ev.call(a, "sheet_read", alias="report", version_id=updated["version_id"])
    ev.check(
        "exact_version_sheet_read",
        [
            {key: read["sheets"]["Report"]["A3"].get(key) for key in ("raw", "value")}
            for read in (old_read, new_read)
        ],
        [
            {"raw": "=A1-A2", "value": INITIAL_MARGIN},
            {"raw": "=SUM(A1,-A2)", "value": DELIVERY_MARGIN},
        ],
    )
    recalculated = ev.call(a, "sheet_recalculate", alias="report")
    valid_bytes = immutable_bytes(world, recalculated)
    ev.check(
        "recalculate_preserves_formula_and_scalar",
        [
            recalculated["object_id"] == first["object_id"],
            recalculated["version_id"] != updated["version_id"],
            workbook_cell(valid_bytes),
        ],
        [True, True, {"raw": "=SUM(A1,-A2)", "cached": DELIVERY_MARGIN, "cache_type": "n"}],
    )
    json_report = ev.call(
        b,
        "create_object",
        alias="explanation",
        filename="report.json",
        kind="json",
        data={"margin": DELIVERY_MARGIN},
        deliverable_role="report",
    )
    ev.check(
        "json_delivery_actual_bytes",
        json.loads(immutable_bytes(world, json_report)),
        {"margin": DELIVERY_MARGIN},
    )
    good = ev.call(a, "submit", work_id="work-1", artifacts=["report"])
    bsub = ev.call(b, "submit", work_id="work-1", artifacts=["explanation"])
    fixed = copy.deepcopy(submission(world, "A", "work-1", good["submission_id"]))
    ev.check(
        "xlsx_submission_pins_current_version",
        fixed["artifact_versions"],
        {first["object_id"]: recalculated["version_id"]},
    )
    valid_eval = ev.evaluate("A", "work-1", good["submission_id"])
    b_eval = ev.evaluate("B", "work-1", bsub["submission_id"])
    ev.check("xlsx_fixed_submission_content_passes", valid_eval["passed"], True)
    ev.check("json_fixed_submission_content_passes", b_eval["passed"], True)
    ev.snapshot("valid-submissions")
    bad = ev.call(a, "sheet_update", alias="report", cells={"Report!A3": "=1/0"})
    bad_bytes = immutable_bytes(world, bad)
    bad_cell = workbook_cell(bad_bytes)
    ev.check(
        "error_formula_persisted_real_bytes",
        [
            bad["version_id"] != recalculated["version_id"],
            bad_cell["raw"],
            bad_cell["cache_type"],
            isinstance(bad_cell["cached"], str) and bad_cell["cached"].startswith("#"),
        ],
        [True, "=1/0", "e", True],
    )
    bad_sub = ev.call(a, "submit", work_id="work-error", artifacts=["report"])
    fixed_bad = submission(world, "A", "work-error", bad_sub["submission_id"])
    ev.check(
        "bad_submission_executes",
        fixed_bad["artifact_versions"],
        {first["object_id"]: bad["version_id"]},
    )
    before_evaluation = world.store.load()
    before_evaluation_files = all_file_hashes(world)
    bad_eval = ev.evaluate("A", "work-error", bad_sub["submission_id"])
    old_eval = ev.evaluate("A", "work-1", good["submission_id"])
    ev.check(
        "independent_evaluation_rejects_formula_error",
        [bad_eval["passed"], any(not check["passed"] for check in bad_eval.get("checks", []))],
        [False, True],
    )
    ev.check(
        "old_submission_and_file_remain_immutable",
        [
            submission(world, "A", "work-1", good["submission_id"]) == fixed,
            immutable_bytes(world, recalculated) == valid_bytes,
        ],
        [True, True],
    )
    ev.check("old_submission_still_passes_after_bad_current", old_eval["passed"], True)
    ev.check(
        "evaluation_reads_only_pinned_versions",
        [[record["exact_read_set"], record["exact_path_set"]] for record in ev.evaluations],
        [[True, True]] * 4,
    )
    ev.check(
        "evaluation_preserves_state_and_file_bytes",
        [
            world.store.load() == before_evaluation,
            all_file_hashes(world) == before_evaluation_files,
        ],
        [True, True],
    )
    ev.check(
        "crossproject_samefilename_remains_independent",
        [
            workbook_cell(immutable_bytes(world, b_book)),
            world.store.load()["artifacts"][b_book["object_id"]]["current_version"],
        ],
        [{"raw": 99, "cached": 99, "cache_type": "n"}, b_book["version_id"]],
    )
    before_rejection, files = world.store.load(), all_file_hashes(world)
    same_project = ev.call(
        world.session("bob", "A"),
        "sheet_update",
        must_succeed=False,
        object_id=first["object_id"],
        cells={"Report!A3": 999},
    )
    cross_project = ev.call(
        b,
        "sheet_update",
        must_succeed=False,
        object_id=first["object_id"],
        cells={"Report!A3": 999},
    )
    ev.check("unauthorized_same_project_edit_rejected", same_project["ok"], False)
    ev.check("unauthorized_cross_project_edit_rejected", cross_project["ok"], False)
    final = world.store.load()
    ev.check(
        "unauthorized_edit_preserves_formal_objects_bytes",
        [final["artifacts"] == before_rejection["artifacts"], all_file_hashes(world) == files],
        [True, True],
    )
    ids = [call["response"]["command_id"] for call in ev.calls]
    commits = final["operation_commits"]
    revisions = [commits[cid]["committed_revision"] for cid in ids]
    ev.check(
        "common_action_journal_records_both_formats",
        [
            len(set(ids)) == len(ids),
            revisions == sorted(set(revisions)),
            all(
                commits[cid]["kind"] == "command"
                and commits[cid]["receipt"]["committed_revision"]
                == commits[cid]["committed_revision"]
                for cid in ids
            ),
            all(call["response"]["command_committed"] for call in ev.calls),
            final["instance_id"] == initial["instance_id"],
        ],
        [True] * 5,
    )
    ev.snapshot("final")


def run_experiment(output):
    output.mkdir(parents=True)
    write_json(output / "protocol.json", PROTOCOL)
    ev = Evidence(output)
    error = None
    try:
        run_n1(ev)
    except Exception as caught:
        error = {"error": f"{type(caught).__name__}: {caught}", "traceback": traceback.format_exc()}
    executed = {record["name"] for record in ev.checks}
    for name in CHECKS:
        if name not in executed:
            ev.checks.append(
                {
                    "name": name,
                    "expected": "See preregistered N1 protocol",
                    "observed": None,
                    "passed": False,
                    "not_executed": True,
                }
            )
    if ev.world:
        ev.snapshot("end-of-run")
    calls = write_json(output / "calls.json", ev.calls)
    evaluations = write_json(output / "evaluations.json", ev.evaluations)
    return {
        "group": "N1",
        "checks": ev.checks,
        "error": error,
        "snapshots": ev.snapshots,
        "calls": calls,
        "evaluations": evaluations,
        "check_count": len(ev.checks),
        "check_pass_count": sum(record["passed"] for record in ev.checks),
        "not_executed_count": sum(record["not_executed"] for record in ev.checks),
        "passed": error is None and all(record["passed"] for record in ev.checks),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Use a fresh output directory")
    before, script_before = code_identity(), digest(Path(__file__).read_bytes())
    started = datetime.now(timezone.utc).isoformat()
    result = run_experiment(args.output)
    after = code_identity()
    report = {
        "suite": PROTOCOL["suite"],
        "protocol": PROTOCOL,
        "started_at": started,
        "ended_at": datetime.now(timezone.utc).isoformat(),
        "code_identity_before": before,
        "code_identity_after": after,
        "script_sha256_before": script_before,
        "script_sha256_after": digest(Path(__file__).read_bytes()),
        "source_tree_unchanged": before["source_tree_sha256"] == after["source_tree_sha256"],
        "result": result,
        "api_calls": 0,
        "gpu_used": False,
        "training_performed": False,
    }
    write_json(args.output / "report.json", report)
    print(
        f"N1: {result['check_pass_count']}/{result['check_count']} checks; "
        f"{result['not_executed_count']} not executed"
    )
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
