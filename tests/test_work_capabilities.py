"""Capability bytes and independent fixed-version evaluation, without a World mock."""

import copy
import hashlib
import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

from proworksim.adapters.capabilities import (
    capability_for,
    capability_tools,
    encode,
    read,
    recalculate_xlsx,
    update_xlsx,
)
from proworksim.domains.work_product import evaluate_submission, validate_content_contract
from proworksim.storage import Store


def ledger(tmp_path, values, contract, submitted=None, adoptions=None):
    store = Store(tmp_path)
    state = {"artifacts": {}, "adoptions": {}, "adoption_view": {}}
    for aid, kind, versions in values:
        artifact = {
            "artifact_id": aid,
            "kind": kind,
            "filename": aid + "." + kind,
            "versions": {},
            "current_version": "v" + str(len(versions)),
            "deliverable_role": "report",
        }
        state["artifacts"][aid] = artifact
        for number, content in enumerate(versions, 1):
            vid = "v" + str(number)
            artifact["versions"][vid] = {"sha256": hashlib.sha256(content).hexdigest()}
            path = store.version_path(artifact, vid)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    item = {"project_id": "B", "work_item_id": "B::work"}
    sub = {
        "submission_id": "sub-1",
        "requirement_snapshot": {"deliverable_contract": contract},
        "artifact_versions": submitted or {values[0][0]: "v1"},
        "review": {"decision": "accepted"},
    }
    if adoptions is not None:
        sub["adoption_snapshot"] = adoptions
    return store, state, item, sub


def test_real_formula_bytes_are_deterministic_and_errors_are_not_repaired():
    source = {"cells": {"Report!A1": 10, "Report!A2": 4, "Report!A3": "=A1-A2"}}
    original = encode("xlsx", source)
    assert original == encode("xlsx", source)
    changed = update_xlsx(original, {"Report!A1": 12, "Report!A3": "=SUM(A1,-A2)"})
    assert changed == update_xlsx(original, {"Report!A1": 12, "Report!A3": "=SUM(A1,-A2)"})
    assert recalculate_xlsx(changed) == changed
    values = load_workbook(io.BytesIO(changed), data_only=True)
    raw = load_workbook(io.BytesIO(changed), data_only=False)
    assert values["Report"]["A3"].value == 8
    assert raw["Report"]["A3"].value == "=SUM(A1,-A2)"
    wrong = update_xlsx(changed, {"Report!A3": "=1/0"})
    assert read("xlsx", wrong)["Report"]["A3"]["error"]
    assert load_workbook(io.BytesIO(wrong), data_only=False)["Report"]["A3"].value == "=1/0"
    assert load_workbook(io.BytesIO(wrong), data_only=True)["Report"]["A3"].data_type == "e"
    assert recalculate_xlsx(wrong) == wrong
    with zipfile.ZipFile(io.BytesIO(wrong)) as archive:
        assert all(entry.date_time == (1980, 1, 1, 0, 0, 0) for entry in archive.infolist())


def test_capability_registry_is_finite_and_schemas_expose_only_enabled_tools():
    assert capability_for("xlsx").application == "spreadsheets"
    assert capability_for("json").suffix == ".json"
    assert capability_tools(["files"]) == []
    assert {tool["name"] for tool in capability_tools(["files", "spreadsheets"])} == {
        "sheet_read",
        "sheet_update",
        "sheet_recalculate",
    }
    with pytest.raises(ValueError):
        capability_for("shell")
    with pytest.raises(ValueError):
        encode("xlsx", {"cells": {}, "run": "anything"})
    with pytest.raises(ValueError):
        update_xlsx(encode("xlsx", {}), {"Report!A1": float("nan")})


def test_evaluator_reads_immutable_submission_and_does_not_call_formula_engine(
    tmp_path, monkeypatch
):
    from proworksim.spreadsheet import Spreadsheet

    good = encode("xlsx", {"Report!A1": 10, "Report!A2": 4, "Report!A3": "=A1-A2"})
    wrong = update_xlsx(good, {"Report!A3": "=1/0"})
    contract = {
        "content_checks": [
            {"kind": "xlsx_cell_equals", "sheet": "Report", "cell": "A3", "expected": 6},
            {"kind": "xlsx_no_formula_errors"},
        ]
    }
    store, state, item, sub = ledger(tmp_path, [("report", "xlsx", [good, wrong])], contract)
    monkeypatch.setattr(Spreadsheet, "value", lambda *_: pytest.fail("Evaluator called engine"))
    monkeypatch.setattr(store, "content", lambda *_: pytest.fail("Evaluator read materialization"))
    before = copy.deepcopy((state, sub))
    result = evaluate_submission(store, state, item, sub)
    assert result["passed"]
    assert result["read_set"] == [
        {
            "object_id": "report",
            "version_id": "v1",
            "sha256": hashlib.sha256(good).hexdigest(),
        }
    ]
    sub["artifact_versions"] = {"report": "v2"}
    result = evaluate_submission(store, state, item, sub)
    assert not result["passed"]
    assert all(not check["passed"] for check in result["checks"])
    assert result["institutional_review"] == {"decision": "accepted"}
    assert state == before[0]


def test_json_contract_accepts_split_deliverable_but_rejects_conflicts_and_boolean(tmp_path):
    contract = {
        "required_fields": ["margin", "note"],
        "content_checks": [
            {"kind": "json_field_equals", "path": "margin", "expected": 6},
        ],
    }
    values = [
        ("summary", "json", [encode("json", {"margin": 6})]),
        ("note", "json", [encode("json", {"note": "The declared interface was used."})]),
        ("conflict", "json", [encode("json", {"margin": 7})]),
        ("boolean", "json", [encode("json", {"margin": True})]),
    ]
    store, state, item, sub = ledger(tmp_path, values, contract, {"summary": "v1", "note": "v1"})
    assert evaluate_submission(store, state, item, sub)["passed"]
    sub["artifact_versions"]["conflict"] = "v1"
    assert evaluate_submission(store, state, item, sub)["conflicting_fields"] == ["margin"]
    sub["artifact_versions"] = {"boolean": "v1", "note": "v1"}
    sub["requirement_snapshot"]["deliverable_contract"]["content_checks"][0]["expected"] = 1
    assert not evaluate_submission(store, state, item, sub)["passed"]


def test_source_interface_checks_snapshot_content_and_keeps_history_stable(tmp_path):
    contract = {
        "content_checks": [
            {
                "kind": "json_matches_source_cell",
                "path": ["margin"],
                "reference_path": ["source_ref"],
                "adoption_alias": "input",
                "sheet": "Report",
                "cell": "A1",
            }
        ]
    }
    a1, a2 = encode("xlsx", {"Report!A1": 6}), encode("xlsx", {"Report!A1": 8})
    source_ref = {"object_id": "source", "version_id": "v1"}
    report1 = encode("json", {"margin": 6, "source_ref": source_ref})
    adoption = {
        "object_id": "source",
        "version_id": "v1",
        "target_version": "v1",
        "policy": "current_published",
        "work_ids": ["B::work"],
    }
    store, state, item, sub = ledger(
        tmp_path,
        [
            ("report", "json", [report1]),
            ("source", "xlsx", [a1, a2]),
        ],
        contract,
        adoptions={"input": adoption},
    )
    assert evaluate_submission(store, state, item, sub)["passed"]
    # Later state changes do not rewrite evaluation of the original fixed submission.
    state["adoptions"]["B::input"] = {**adoption, "version_id": "v2"}
    state["adoption_view"]["B::input"] = {"target_version": "v2"}
    assert evaluate_submission(store, state, item, sub)["passed"]
    # A newly submitted stale source, or a new label attached to the old result, fails.
    adoption["target_version"] = "v2"
    sub["adoption_snapshot"]["input"] = copy.deepcopy(adoption)
    assert not evaluate_submission(store, state, item, sub)["passed"]
    adoption["version_id"] = "v2"
    sub["adoption_snapshot"]["input"] = copy.deepcopy(adoption)
    wrong = encode("json", {"margin": 6, "source_ref": {**source_ref, "version_id": "v2"}})
    path = store.version_path(state["artifacts"]["report"], "v1")
    path.write_bytes(wrong)
    state["artifacts"]["report"]["versions"]["v1"]["sha256"] = hashlib.sha256(wrong).hexdigest()
    result = evaluate_submission(store, state, item, sub)
    assert not result["passed"] and result["checks"][0]["expected"] == 8
    del sub["adoption_snapshot"]
    result = evaluate_submission(store, state, item, sub)
    assert not result["passed"] and "UNASSESSED" in result["checks"][0]["reason"]


def test_evaluator_rejects_uncommitted_byte_change_and_unknown_check(tmp_path):
    store, state, item, sub = ledger(tmp_path, [("report", "json", [b'{"margin": 6}'])], {})
    store.version_path(state["artifacts"]["report"], "v1").write_bytes(
        json.dumps({"margin": 8}).encode()
    )
    assert "digest" in evaluate_submission(store, state, item, sub)["errors"][0]
    with pytest.raises(ValueError):
        validate_content_contract({"content_checks": [{"kind": "python_eval"}]})
    with pytest.raises(ValueError):
        validate_content_contract({"content_checks": [{"kind": "json_field_equals", "path": "x"}]})
