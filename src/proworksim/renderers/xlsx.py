from openpyxl import Workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter


def relocate(book, layout):
    if layout.layout_id == "standard":
        return book
    result = Workbook()
    result.remove(result.active)
    for source in book:
        name = layout.input_sheet if source.title == "Inputs" else layout.output_sheet
        target = result.create_sheet(name)
        offset = layout.input_offset if source.title == "Inputs" else layout.output_offset
        value_column = layout.input_column if source.title == "Inputs" else layout.output_column
        column_delta = column_index_from_string(value_column) - 2
        for row in source:
            for cell in row:
                if cell.value is None:
                    continue
                value = cell.value
                if cell.data_type == "f":
                    fragments = []
                    for token in Tokenizer(value).items:
                        if token.type == "OPERAND" and token.subtype == "RANGE":
                            reference = (
                                token.value
                                if "!" in token.value
                                else f"{source.title}!{token.value}"
                            )
                            fragments.append(layout.address(reference.replace("$", "")))
                        else:
                            fragments.append(token.value)
                    value = "=" + "".join(fragments)
                address = f"{get_column_letter(cell.column + column_delta)}{cell.row + offset}"
                target[address] = value
        target.column_dimensions[get_column_letter(column_delta + 1)].width = 36
        target.column_dimensions[value_column].width = 24
    return result
