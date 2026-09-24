"""Read-only content contracts for finite JSON/XLSX work products.

The evaluator reads immutable submitted bytes and persisted XLSX value caches.
It does not invoke Spreadsheet or infer correctness from an institutional review,
adoption status label, current materialization, or a hidden answer workbook.
Scalar expected values must be declared in the contract by its author.
"""

import copy
import hashlib
import io
import json
import math

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from ..core.adoption import binding_key, require_version
from ..evaluation import EvaluationInputError, combined_status
from ..core.references import VersionRef
from . import reconciliation, research_review, executable_project

DOMAIN_CHECKS = {module.CHECK_KIND: module for module in (reconciliation, research_review, executable_project)}


_KINDS = {
    "json_field_equals",
    "xlsx_cell_equals",
    "xlsx_no_formula_errors",
    "json_matches_source_cell",
    "json_matches_source_field",
    "json_linear_sources",
}
_COMMON = {"kind", "role"}
EVALUATOR_VERSION = "finite-products-v0.11"


def _path(value, label):
    if isinstance(value, str):
        value = value.split(".")
    if (
        not isinstance(value, list)
        or not value
        or any(not isinstance(part, str) or not part for part in value)
    ):
        raise ValueError(label + " must be a nonempty JSON field path")
    return value


def _cell(check):
    if not isinstance(check.get("sheet"), str) or not check["sheet"]:
        raise ValueError("XLSX check requires a sheet name")
    cell = check.get("cell")
    if not isinstance(cell, str):
        raise ValueError("XLSX check requires a cell address")
    try:
        row, column = coordinate_to_tuple(cell)
    except (ValueError, KeyError) as exc:
        raise ValueError("Invalid XLSX cell address") from exc
    if not 1 <= row <= 2000 or not 1 <= column <= 100:
        raise ValueError("XLSX contract cell exceeds bounded capability")
    check["cell"] = cell.upper()


def _finite_number(value):
    return type(value) is int or (type(value) is float and math.isfinite(value))


def _linear_sources(check):
    sources = check.get("sources")
    if not isinstance(sources, list) or not 2 <= len(sources) <= 8:
        raise ValueError("Linear source contract requires two to eight declared inputs")
    aliases = set()
    for source in sources:
        if not isinstance(source, dict) or source.get("kind") not in {"json_field", "xlsx_cell"}:
            raise ValueError("Linear source must select a JSON field or XLSX cell")
        fields = {"alias", "kind", "reference_path", "coefficient"}
        alias = source.get("alias")
        if not isinstance(alias, str) or not alias or alias in aliases:
            raise ValueError("Linear source aliases must be distinct nonempty strings")
        aliases.add(alias)
        source["reference_path"] = _path(source.get("reference_path"), "reference_path")
        source.setdefault("coefficient", 1)
        if not _finite_number(source["coefficient"]):
            raise ValueError("Linear source coefficient must be a finite number")
        if source["kind"] == "json_field":
            fields.add("source_path")
            source["source_path"] = _path(source.get("source_path"), "source_path")
        else:
            fields |= {"sheet", "cell"}
            _cell(source)
        if set(source) - fields:
            raise ValueError("Unknown linear source fields")
    check.setdefault("constant", 0)
    if not _finite_number(check["constant"]):
        raise ValueError("Linear source constant must be a finite number")


def validate_content_contract(contract):
    """Validate only this module's extension and preserve the base contract."""
    if not isinstance(contract, dict):
        raise ValueError("Deliverable contract must be an object")
    result = copy.deepcopy(contract)
    checks = result.setdefault("content_checks", [])
    if not isinstance(checks, list) or len(checks) > 100:
        raise ValueError("content_checks must contain at most 100 finite checks")
    for index, check in enumerate(checks):
        if isinstance(check, dict) and check.get("kind") in DOMAIN_CHECKS:
            checks[index] = DOMAIN_CHECKS[check["kind"]].validate_check(check)
            continue
        if not isinstance(check, dict) or check.get("kind") not in _KINDS:
            raise ValueError("Unknown work product content check")
        kind = check["kind"]
        fields = set(_COMMON)
        if kind.startswith("json_"):
            fields.add("path")
            check["path"] = _path(check.get("path"), "path")
        if kind in {"xlsx_cell_equals", "json_matches_source_cell"}:
            fields |= {"sheet", "cell"}
            _cell(check)
        if kind.endswith("_equals"):
            fields.add("expected")
            expected = check.get("expected")
            if (
                "expected" not in check
                or type(expected) not in (str, int, float, bool, type(None))
                or (type(expected) in (int, float) and not math.isfinite(expected))
            ):
                raise ValueError("Expected content must be an explicit finite JSON scalar")
        if kind.startswith("json_matches_source_"):
            fields |= {"reference_path", "adoption_alias"}
            check["reference_path"] = _path(check.get("reference_path"), "reference_path")
            if not isinstance(check.get("adoption_alias"), str) or not check["adoption_alias"]:
                raise ValueError("Source check requires an adoption_alias")
            if kind == "json_matches_source_field":
                fields.add("source_path")
                check["source_path"] = _path(check.get("source_path"), "source_path")
        if kind == "json_linear_sources":
            fields |= {"sources", "constant"}
            _linear_sources(check)
        if "role" in check and (not isinstance(check["role"], str) or not check["role"]):
            raise ValueError("Content check role must be a nonempty string")
        if set(check) - fields:
            raise ValueError(
                "Unknown content check fields: " + ", ".join(sorted(set(check) - fields))
            )
    return result


def _at_path(data, path):
    for key in path:
        if not isinstance(data, dict) or key not in data:
            raise ValueError("Missing JSON path: " + ".".join(path))
        data = data[key]
    return data


def _equal(actual, expected):
    """JSON equality at every depth, preserving the finite numeric contract."""
    if type(actual) in (int, float) and type(expected) in (int, float):
        return (
            (type(actual) is int or math.isfinite(actual))
            and (type(expected) is int or math.isfinite(expected))
            and actual == expected
        )
    if type(actual) is not type(expected):
        return False
    if isinstance(actual, dict):
        return actual.keys() == expected.keys() and all(
            _equal(value, expected[key]) for key, value in actual.items()
        )
    if isinstance(actual, list):
        return len(actual) == len(expected) and all(
            _equal(left, right) for left, right in zip(actual, expected)
        )
    return actual == expected


def _merge_json(documents):
    """Union whole top-level fields; conflicting values have no selected winner.

    A conflict remains a conflict when a third document repeats either value.
    This deliberately does not recursively merge overlapping nested objects.
    """
    values, conflicts = {}, set()
    for document in documents:
        for key, value in document.items():
            if key in values and not _equal(values[key], value):
                conflicts.add(key)
            else:
                values[key] = value
    return {key: value for key, value in values.items() if key not in conflicts}, conflicts


def _books(content):
    return (
        load_workbook(io.BytesIO(content), data_only=False),
        load_workbook(io.BytesIO(content), data_only=True),
    )


def _cached_cell(content, sheet, cell):
    raw, values = _books(content)
    if sheet not in raw.sheetnames:
        raise ValueError("Missing XLSX sheet: " + sheet)
    original, cached = raw[sheet][cell], values[sheet][cell]
    if original.data_type == "e" or cached.data_type == "e":
        raise ValueError("XLSX cell has persisted error: " + str(cached.value))
    if original.data_type == "f" and cached.value is None:
        raise ValueError("XLSX formula has no persisted value cache")
    return cached.value


def evaluate_submission(store, state, item, submission, *, _context=None):
    """Evaluate fixed submitted versions; return all checks, exact reads and errors.

    Source-linked checks require the submission's adoption_snapshot, including
    its policy target at submission time. Later releases or adoption changes do
    not rewrite the meaning of a historical submission.
    """
    contract = validate_content_contract(
        submission["requirement_snapshot"].get("deliverable_contract", {})
    )
    reads, content_cache, errors, input_statuses = [], {}, [], []
    if _context is not None:
        _context.update(read_set=reads, errors=errors)

    def read_version(aid, vid):
        key = (aid, vid)
        if key not in content_cache:
            artifact = state["artifacts"].get(aid)
            if artifact is None or vid not in artifact["versions"]:
                raise EvaluationInputError("source_unavailable", "Unknown exact content version")
            try:
                content = store.version_path(artifact, vid).read_bytes()
            except OSError as exc:
                raise EvaluationInputError("source_unavailable", str(exc)) from exc
            sha = hashlib.sha256(content).hexdigest()
            if sha != artifact["versions"][vid]["sha256"]:
                raise EvaluationInputError(
                    "source_unavailable", "Immutable content does not match its committed digest"
                )
            content_cache[key] = content
            reads.append({"object_id": aid, "version_id": vid, "sha256": sha})
        return content_cache[key]

    submitted = []
    for aid, vid in submission["artifact_versions"].items():
        artifact = state["artifacts"][aid]
        try:
            content = read_version(aid, vid)
            data = json.loads(content) if artifact["kind"] == "json" else None
            if artifact["kind"] == "json" and not isinstance(data, dict):
                raise EvaluationInputError("structure_failure", "Submitted JSON must be an object")
            submitted.append(
                {"artifact": artifact, "version_id": vid, "content": content, "data": data}
            )
        except ValueError as exc:
            errors.append(str(exc))
            input_statuses.append(getattr(exc, "status", "structure_failure"))
    combined, conflicts = _merge_json(
        entry["data"] for entry in submitted if entry["data"] is not None
    )
    missing = sorted(set(contract.get("required_fields", [])) - (set(combined) | conflicts))

    def bound_source(selected, json_data, value_path, reference_path, alias):
        reference = _at_path(json_data, reference_path)
        if not isinstance(reference, dict):
            raise ValueError("Source reference must be an exact object/version mapping")
        ref = VersionRef.from_mapping(reference)
        key = binding_key(item["work_item_id"], alias)
        snapshots = submission.get("adoption_snapshot", {})
        adoption = snapshots.get(key)
        if adoption is None:
            raise EvaluationInputError("unassessed", "Submission has no adoption snapshot")
        if (adoption["object_id"], adoption["version_id"]) != (
            ref.object_id,
            ref.version_id,
        ):
            raise EvaluationInputError(
                "content_failure", "Content source does not match the adopted exact version"
            )
        if (
            adoption.get("work_id") != item["work_item_id"]
            or adoption.get("work_ids") != [item["work_item_id"]]
            or adoption.get("project_id") != item["project_id"]
            or adoption.get("alias") != alias
            or adoption.get("requirement_version") != submission["requirement_version"]
        ):
            raise EvaluationInputError(
                "content_failure", "Adoption does not cover the evaluated work edition"
            )
        require_version(
            submission["requirement_snapshot"],
            alias,
            adoption["version_id"],
            adoption["policy"],
        )
        if adoption["policy"] != "fixed":
            target = adoption.get("target_version")
            if target is None or ref.version_id != target:
                raise EvaluationInputError(
                    "content_failure",
                    "Content source does not match the policy target at submission",
                )
        binding_files = []
        for entry in selected:
            if entry["data"] is None:
                continue
            contributes = False
            for path in (value_path, reference_path):
                try:
                    _at_path(entry["data"], path)
                    contributes = True
                except ValueError:
                    pass
            if not contributes:
                continue
            metadata = entry["artifact"]["versions"][entry["version_id"]]
            dependencies = [
                VersionRef.from_mapping(value) for value in metadata.get("derived_from", [])
            ]
            if ref not in dependencies:
                raise EvaluationInputError(
                    "content_failure",
                    "Contributing submitted file lacks the exact source dependency: "
                    + entry["artifact"]["artifact_id"]
                    + "@"
                    + entry["version_id"],
                )
            binding_files.append(
                {
                    "object_id": entry["artifact"]["artifact_id"],
                    "version_id": entry["version_id"],
                }
            )
        return ref, binding_files

    checks = []
    if _context is not None:
        _context["checks"] = checks
    for spec in contract["content_checks"]:
        if _context is not None:
            _context["active_check"] = copy.deepcopy(spec)
        outcome = {"kind": spec["kind"], "contract": copy.deepcopy(spec), "passed": False}
        try:
            selected = [
                entry
                for entry in submitted
                if ("role" not in spec or entry["artifact"].get("deliverable_role") == spec["role"])
            ]
            json_data, selected_conflicts = _merge_json(
                entry["data"] for entry in selected if entry["data"] is not None
            )
            kind = spec["kind"]
            if kind.startswith("json_") or kind in DOMAIN_CHECKS:
                required_roots = {spec["path"][0]}
                if "reference_path" in spec:
                    required_roots.add(spec["reference_path"][0])
                if kind == "json_linear_sources" or kind in DOMAIN_CHECKS:
                    required_roots.update(source["reference_path"][0] for source in spec["sources"])
                relevant_conflicts = sorted(required_roots & selected_conflicts)
                if relevant_conflicts:
                    raise ValueError("Conflicting JSON fields: " + ", ".join(relevant_conflicts))
            if kind in DOMAIN_CHECKS:
                if kind == executable_project.CHECK_KIND:
                    executable_project.validate_execution_evidence(state, item, submission, selected, spec)
                source_data, exact_sources = {}, []
                for source_spec in spec["sources"]:
                    ref, files = bound_source(
                        selected,
                        json_data,
                        spec["path"],
                        source_spec["reference_path"],
                        source_spec["alias"],
                    )
                    if state["artifacts"][ref.object_id]["kind"] != "json":
                        raise ValueError("This finite domain contract requires JSON inputs")
                    try:
                        source_data[source_spec["alias"]] = json.loads(
                            read_version(ref.object_id, ref.version_id)
                        )
                    except ValueError as exc:
                        raise EvaluationInputError("source_unavailable", str(exc)) from exc
                    exact_sources.append(
                        {
                            "alias": source_spec["alias"],
                            "reference": ref.to_dict(),
                            "binding_files": files,
                        }
                    )
                outcome.update(
                    DOMAIN_CHECKS[kind].evaluate_check(spec, json_data, source_data.__getitem__)
                )
                outcome["sources"] = exact_sources
            elif kind == "json_field_equals":
                actual, expected = _at_path(json_data, spec["path"]), spec["expected"]
                outcome.update(actual=actual, expected=expected, passed=_equal(actual, expected))
            elif kind == "xlsx_cell_equals":
                candidates = [entry for entry in selected if entry["artifact"]["kind"] == "xlsx"]
                if len(candidates) != 1:
                    raise ValueError("XLSX cell contract requires exactly one matching workbook")
                actual = _cached_cell(candidates[0]["content"], spec["sheet"], spec["cell"])
                outcome.update(
                    actual=actual,
                    expected=spec["expected"],
                    passed=_equal(actual, spec["expected"]),
                )
            elif kind == "xlsx_no_formula_errors":
                candidates = [entry for entry in selected if entry["artifact"]["kind"] == "xlsx"]
                if not candidates:
                    raise ValueError("No submitted workbook for formula check")
                formula_errors = []
                for entry in candidates:
                    raw, values = _books(entry["content"])
                    for sheet in raw:
                        for row in sheet:
                            for cell in row:
                                cached = values[sheet.title][cell.coordinate]
                                if (
                                    cell.data_type == "e"
                                    or cached.data_type == "e"
                                    or (cell.data_type == "f" and cached.value is None)
                                ):
                                    formula_errors.append(
                                        {
                                            "object_id": entry["artifact"]["artifact_id"],
                                            "cell": sheet.title + "!" + cell.coordinate,
                                            "formula": cell.value,
                                            "cached": cached.value,
                                        }
                                    )
                outcome.update(formula_errors=formula_errors, passed=not formula_errors)
            elif kind == "json_linear_sources":
                expected = spec["constant"]
                terms = []
                for source_spec in spec["sources"]:
                    ref, binding_files = bound_source(
                        selected,
                        json_data,
                        spec["path"],
                        source_spec["reference_path"],
                        source_spec["alias"],
                    )
                    content = read_version(ref.object_id, ref.version_id)
                    artifact = state["artifacts"][ref.object_id]
                    if source_spec["kind"] == "json_field":
                        if artifact["kind"] != "json":
                            raise ValueError("Linear JSON source must be a JSON object")
                        value = _at_path(json.loads(content), source_spec["source_path"])
                    else:
                        if artifact["kind"] != "xlsx":
                            raise ValueError("Linear XLSX source must be a workbook")
                        value = _cached_cell(content, source_spec["sheet"], source_spec["cell"])
                    if not _finite_number(value):
                        raise ValueError("Linear source value must be a finite number")
                    contribution = source_spec["coefficient"] * value
                    expected += contribution
                    if not _finite_number(contribution) or not _finite_number(expected):
                        raise ValueError("Linear source result must remain finite")
                    terms.append(
                        {
                            "alias": source_spec["alias"],
                            "source": ref.to_dict(),
                            "value": value,
                            "coefficient": source_spec["coefficient"],
                            "contribution": contribution,
                            "binding_files": binding_files,
                        }
                    )
                actual = _at_path(json_data, spec["path"])
                outcome.update(
                    actual=actual,
                    expected=expected,
                    sources=terms,
                    passed=_finite_number(actual) and actual == expected,
                )
            else:
                ref, binding_files = bound_source(
                    selected,
                    json_data,
                    spec["path"],
                    spec["reference_path"],
                    spec["adoption_alias"],
                )
                outcome["binding_files"] = binding_files
                source = read_version(ref.object_id, ref.version_id)
                artifact = state["artifacts"][ref.object_id]
                if kind == "json_matches_source_cell":
                    if artifact["kind"] != "xlsx":
                        raise ValueError("Adopted source must be an XLSX workbook")
                    expected = _cached_cell(source, spec["sheet"], spec["cell"])
                else:
                    if artifact["kind"] != "json":
                        raise ValueError("Adopted source must be JSON")
                    expected = _at_path(json.loads(source), spec["source_path"])
                actual = _at_path(json_data, spec["path"])
                outcome.update(
                    actual=actual,
                    expected=expected,
                    source=ref.to_dict(),
                    passed=_equal(actual, expected),
                )
            if not outcome["passed"]:
                outcome.setdefault(
                    "reason", "Submitted content does not satisfy the declared contract"
                )
        except ValueError as exc:
            outcome["reason"] = str(exc)
            outcome["status"] = getattr(exc, "status", "structure_failure")
        outcome.setdefault("status", "pass" if outcome["passed"] else "content_failure")
        checks.append(outcome)
    statuses = input_statuses + [check["status"] for check in checks]
    if missing or conflicts:
        statuses.append("structure_failure")
    if not contract["content_checks"]:
        statuses.append("unassessed")
    status = combined_status(statuses)
    return {
        "submission_id": submission["submission_id"],
        "evaluator_version": EVALUATOR_VERSION,
        "passed": status == "pass",
        "status": status,
        "missing_fields": missing,
        "conflicting_fields": sorted(set(conflicts)),
        "checks": checks,
        "errors": errors,
        "read_set": reads,
        "scope": "finite_managed_content_contract",
        "institutional_review": copy.deepcopy(submission.get("review")),
    }
